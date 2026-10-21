#!/usr/bin/env python3
"""Redraw the MCP Protocol Architecture explainer as NATIVE pptx shapes
matching the deck theme (navy/cyan slate palette, Noto Sans) — not a mermaid image.

Usage:  python3 scripts/redraw_arch_diagram.py ["Training Process.pptx"]

Finds the slide whose title == TITLE, removes any picture + old caption on it,
and draws the Host / Client / Server + two-layer architecture with native shapes.
Idempotent-ish: re-running removes the prior picture but appends fresh shapes, so
run against a clean deck (restore from Training Process.backup.pptx if needed).

Palette sampled from the deck's own slides (ppt/slides/*.xml):
  navy #0F1B2C · cyan #00AEEF / #38BDF8 · border #E2E8F0 · muted #64748B
  off-white #F8FAFC · tint #F1F5F9 · green #34D399 · font Noto Sans
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

NAVY=RGBColor(0x0F,0x1B,0x2C); NAVY2=RGBColor(0x1E,0x2D,0x4A)
CYAN=RGBColor(0x00,0xAE,0xEF); CYAN_LT=RGBColor(0x38,0xBD,0xF8)
WHITE=RGBColor(0xFF,0xFF,0xFF); OFF=RGBColor(0xF8,0xFA,0xFC); TINT=RGBColor(0xF1,0xF5,0xF9)
BORDER=RGBColor(0xE2,0xE8,0xF0); SLATE=RGBColor(0x33,0x41,0x55); MUTED=RGBColor(0x64,0x74,0x8B)
GREEN=RGBColor(0x34,0xD3,0x99); CLOUD=RGBColor(0xCB,0xD5,0xE1); ICE=RGBColor(0xEA,0xF7,0xFD)
FONT="Noto Sans"
TITLE="MCP Protocol Architecture (Host / Client / Server + layers)"


def main(path):
    prs=Presentation(path)
    slide=None
    for s in prs.slides:
        try:
            if s.shapes.title and s.shapes.title.text.strip().split("\n")[0]==TITLE:
                slide=s; break
        except Exception:
            pass
    assert slide, f"slide titled {TITLE!r} not found"
    title_ph=slide.shapes.title

    for sh in list(slide.shapes):
        if sh is title_ph:
            continue
        if sh.shape_type==13:  # picture
            sh._element.getparent().remove(sh._element)
        elif sh.has_text_frame and sh.text_frame.text.strip().startswith("The Host runs"):
            sh._element.getparent().remove(sh._element)  # old caption

    def para(p,text,size,color,bold=False,italic=False,align=PP_ALIGN.CENTER):
        p.alignment=align
        r=p.add_run(); r.text=text
        f=r.font; f.name=FONT; f.size=Pt(size); f.bold=bold; f.italic=italic; f.color.rgb=color

    def box(x,y,w,h,fill,line=None,line_w=1.0,radius=0.08):
        sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        try: sh.adjustments[0]=radius
        except Exception: pass
        if fill is None: sh.fill.background()
        else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
        if line is None: sh.line.fill.background()
        else: sh.line.color.rgb=line; sh.line.width=Pt(line_w)
        sh.shadow.inherit=False
        return sh

    def label(x,y,w,h,lines,anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER):
        tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
        tf.margin_left=Pt(4); tf.margin_right=Pt(4); tf.margin_top=Pt(2); tf.margin_bottom=Pt(2)
        for i,(t,sz,c,b,it) in enumerate(lines):
            para(tf.paragraphs[0] if i==0 else tf.add_paragraph(),t,sz,c,b,it,align)
        return tb

    def text_in(shape,lines,anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER):
        tf=shape.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
        tf.margin_left=Pt(6); tf.margin_right=Pt(6); tf.margin_top=Pt(3); tf.margin_bottom=Pt(3)
        for i,(t,sz,c,b,it) in enumerate(lines):
            para(tf.paragraphs[0] if i==0 else tf.add_paragraph(),t,sz,c,b,it,align)

    def connector(x1,y1,x2,y2,color=CYAN,width=2.25,dashed=False,arrow=True):
        cn=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        cn.line.color.rgb=color; cn.line.width=Pt(width)
        ln=cn.line._get_or_add_ln()
        if dashed: ln.append(ln.makeelement(qn('a:prstDash'),{'val':'sysDash'}))
        if arrow: ln.append(ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}))
        cn.shadow.inherit=False
        return cn

    L=0.7; R=19.3; W=R-L
    for p in title_ph.text_frame.paragraphs:
        for r in p.runs:
            r.font.name=FONT; r.font.size=Pt(26); r.font.bold=True; r.font.color.rgb=NAVY

    # Row 1 — HOST
    hy=1.65; hh=1.95; box(L,hy,W,hh,NAVY,radius=0.05)
    label(L+0.25,hy+0.12,W-0.5,0.4,[("HOST — AI application  (LibreChat · Claude Desktop · Cursor)",13,CYAN,True,False)],MSO_ANCHOR.TOP,PP_ALIGN.LEFT)
    py=hy+0.72; ph=0.95
    ca=box(1.15,py,3.3,ph,OFF,CYAN,1.25,0.16); text_in(ca,[("MCP Client A",12.5,NAVY,True,False),("1:1 connector",9.5,MUTED,False,False)])
    llm=box(8.0,py,4.0,ph,CYAN,radius=0.16); text_in(llm,[("LLM",13,WHITE,True,False),("reasons · picks tools",9.5,ICE,False,False)])
    cb=box(15.55,py,3.3,ph,OFF,CYAN,1.25,0.16); text_in(cb,[("MCP Client B",12.5,NAVY,True,False),("1:1 connector",9.5,MUTED,False,False)])
    connector(4.45,py+ph/2,8.0,py+ph/2,CYAN_LT,1.25,arrow=False)
    connector(12.0,py+ph/2,15.55,py+ph/2,CYAN_LT,1.25,arrow=False)

    # Row 2 — servers
    sy=4.35; shh=2.75
    s1x=L; s1w=11.0; box(s1x,sy,s1w,shh,WHITE,BORDER,1.5,0.045)
    h1=box(s1x,sy,s1w,0.62,NAVY,radius=0.045); text_in(h1,[("SERVER 1 — Router MCP server (ours)",12.5,WHITE,True,False)])
    chips=[("Tools","get_interfaces · configure_interface","controlled by the MODEL",CYAN),
           ("Resources","router://running-config","controlled by the APP / USER",CYAN_LT),
           ("Prompts","/troubleshoot_interface","controlled by the USER",GREEN)]
    cx=s1x+0.35; cw=(s1w-0.7-0.6)/3; cy=sy+0.85; ch=1.65
    for i,(nm,ex,ctrl,accent) in enumerate(chips):
        x=cx+i*(cw+0.3)
        box(x,cy,cw,ch,TINT,BORDER,1.0,0.09)
        box(x+0.18,cy+0.2,0.22,0.22,accent,radius=0.5)
        label(x+0.5,cy+0.14,cw-0.6,0.4,[(nm,12.5,NAVY,True,False)],MSO_ANCHOR.TOP,PP_ALIGN.LEFT)
        label(x+0.18,cy+0.62,cw-0.36,0.7,[(ex,9.5,SLATE,False,False)],MSO_ANCHOR.TOP,PP_ALIGN.LEFT)
        label(x+0.18,cy+ch-0.42,cw-0.36,0.35,[(ctrl,8.5,MUTED,False,True)],MSO_ANCHOR.BOTTOM,PP_ALIGN.LEFT)
    s2x=12.2; s2w=R-s2x; box(s2x,sy,s2w,shh,WHITE,BORDER,1.5,0.045)
    h2=box(s2x,sy,s2w,0.62,NAVY2,radius=0.045); text_in(h2,[("SERVER 2 — another MCP server",12,WHITE,True,False)])
    in2=box(s2x+0.4,sy+0.95,s2w-0.8,1.5,TINT,BORDER,1.0,0.07)
    text_in(in2,[("Tools · Resources · Prompts",12,NAVY,True,False),("same protocol, built once",9.5,MUTED,False,True)])
    connector(2.8,py+ph,6.2,sy,CYAN,2.5); label(3.2,(py+ph+sy)/2-0.2,2.4,0.35,[("MCP · 1:1",10,CYAN,True,False)])
    connector(17.2,py+ph,15.7,sy,CYAN,2.5); label(16.0,(py+ph+sy)/2-0.2,2.4,0.35,[("MCP · 1:1",10,CYAN,True,False)])

    # Row 3 — two-layer band
    ly=7.5; lh=1.9; box(L,ly,W,lh,TINT,BORDER,1.25,0.05)
    label(L+0.25,ly+0.1,W-0.5,0.35,[("Every Client ↔ Server link runs on two layers",11.5,NAVY,True,False)],MSO_ANCHOR.TOP,PP_ALIGN.LEFT)
    by=ly+0.62; bh=1.05
    data=box(L+0.35,by,W/2-0.6,bh,NAVY,radius=0.07)
    text_in(data,[("DATA LAYER — JSON-RPC 2.0",12,CYAN,True,False),("lifecycle · tools / resources / prompts · notifications",9.5,CLOUD,False,False)])
    trans=box(L+W/2+0.25,by,W/2-0.6,bh,CYAN,radius=0.07)
    text_in(trans,[("TRANSPORT LAYER",12,WHITE,True,False),("stdio (local subprocess) · Streamable HTTP (remote: POST + SSE)",9.5,ICE,False,False)])
    connector(6.2,sy+shh,6.2,ly,CYAN,1.5,dashed=True); label(6.4,(sy+shh+ly)/2-0.18,2.0,0.3,[("runs on",9.5,MUTED,False,True)],align=PP_ALIGN.LEFT)

    # footnote
    fy=9.62; foot=box(L,fy,W,0.62,OFF,BORDER,1.0,0.09)
    text_in(foot,[("In the lab:  LibreChat  →  Streamable HTTP  →  :8000/mcp (Bearer)  →  MCP server  →  SSH  →  Router1 (172.18.2.6)",10.5,SLATE,False,False)])

    prs.save(path)
    print("redrew architecture diagram (native shapes) in", path)


if __name__=="__main__":
    main(sys.argv[1] if len(sys.argv)>1 else "Training Process.pptx")
