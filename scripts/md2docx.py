"""Minimal markdown -> docx converter for the workshop walkthrough doc.

Handles: # ## ### headings, paragraphs with **bold** and `code` inline,
fenced code blocks, - bullets, numbered list lines, | tables |, --- rules.
"""

import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

CODE_BG = "F2F2F2"
CODE_COLOR = RGBColor(0x1F, 0x1F, 0x1F)
MONO = "Consolas"


def shade(element, fill):
    shd = element.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): fill})
    element.get_or_add_pPr().append(shd)


def add_inline(par, text, mono_size=10):
    """Render **bold** and `code` spans inside a paragraph."""
    for token in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = par.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = par.add_run(token[1:-1])
            run.font.name = MONO
            run.font.size = Pt(mono_size)
        else:
            par.add_run(token)


def add_code_block(doc, lines):
    for line in lines:
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.left_indent = Inches(0.25)
        shade(par._p, CODE_BG)
        run = par.add_run(line if line else " ")
        run.font.name = MONO
        run.font.size = Pt(9)
        run.font.color.rgb = CODE_COLOR


def add_table(doc, rows):
    header, body = rows[0], [r for r in rows[1:] if not re.match(r"^[\s\-|:]+$", "|".join(r))]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Light Grid Accent 1"
    for col, text in enumerate(header):
        par = table.rows[0].cells[col].paragraphs[0]
        add_inline(par, text, mono_size=9)
        for run in par.runs:
            run.bold = True
    for i, row in enumerate(body, start=1):
        for col, text in enumerate(row):
            if col < len(table.rows[i].cells):
                add_inline(table.rows[i].cells[col].paragraphs[0], text, mono_size=9)
    doc.add_paragraph()


def split_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def main(src, dst):
    text = open(src).read()
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = text.splitlines()
    i, first_h1 = 0, True
    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            j = i + 1
            block = []
            while j < len(lines) and not lines[j].startswith("```"):
                block.append(lines[j])
                j += 1
            add_code_block(doc, block)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            i = j + 1
            continue

        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        if line.startswith("# "):
            heading = doc.add_heading(line[2:].strip(), level=0 if first_h1 else 1)
            if first_h1:
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_h1 = False
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.strip() == "---":
            pass
        elif line.startswith("- "):
            add_inline(doc.add_paragraph(style="List Bullet"), line[2:].strip())
        elif re.match(r"^\d+\. ", line):
            number, _, rest = line.partition(". ")
            add_inline(doc.add_paragraph(style="List Number"), rest.strip())
        elif line.strip():
            add_inline(doc.add_paragraph(), line.strip())
        i += 1

    doc.save(dst)
    print(f"Wrote {dst}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
