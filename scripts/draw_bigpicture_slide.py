#!/usr/bin/env python3
"""Add a SIMPLE high-level 'big picture' slide (native pptx shapes, deck theme)
grounded in the workshop lab topology (mcp-workshop-topo.png):

  MCP-Client desktop (172.18.0.2)              MCP-Server desktop (172.18.1.2)      Router1 (172.18.2.6)
    Host app: LibreChat   --- MCP ---> [lab net] ---> MCP server script (FastMCP)  --- SSH --->  Cisco IOS
      -> MCP Client         Streamable HTTP            -> Tools (get_interfaces …)     Netmiko

One client, one server, one router. Inserted just before the detailed
'Workshop lab topology' explainer slide.

Usage:  python3 scripts/draw_bigpicture_slide.py ["Training Process.pptx"]
Re-running ADDS another copy — run against a clean deck (restore from
Training Process.backup.pptx) or delete the prior slide first.

Palette (sampled from deck): navy #0F1B2C · cyan #00AEEF · green #34D399 ·
border #E2E8F0 · muted #64748B · tint #F1F5F9 · font Noto Sans.
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

NAVY=RGBColor(0x0F,0x1B,0x2C); NAVY2=RGBColor(0x1E,0x2D,0x4A)
CYAN=RGBColor(0x00,0xAE,0xEF); CYAN_LT=RGBColor(0x38,0xBD,0xF8)
WHITE=RGBColor(0xFF,0xFF,0xFF); OFF=RGBColor(0xF8,0xFA,0xFC); TINT=RGBColor(0xF1,0xF5,0xF9)
BORDER=RGBColor(0xE2,0xE8,0xF0); SLATE=RGBColor(0x33,0x41,0x55); MUTED=RGBColor(0x64,0x74,0x8B)
GREEN=RGBColor(0x34,0xD3,0x99); GREEN_D=RGBColor(0x0E,0x8A,0x5E); ICE=RGBColor(0xEA,0xF7,0xFD)
FONT="Noto Sans"
TITLE="The Big Picture — Client · Server · Router"


def main(path):
    prs=Presentation(path)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.title.text=TITLE
    for p in slide.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            r.font.name=FONT; r.font.size=Pt(26); r.font.bold=True; r.font.color.rgb=NAVY

    def para(p,t,sz,c,b=False,it=False,al=PP_ALIGN.CENTER):
        p.alignment=al; r=p.add_run(); r.text=t
        f=r.font; f.name=FONT; f.size=Pt(sz); f.bold=b; f.italic=it; f.color.rgb=c
    def box(x,y,w,h,fill,line=None,lw=1.0,rad=0.08,shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        sh=slide.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
        try:
            if shape==MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0]=rad
        except Exception: pass
        if fill is None: sh.fill.background()
        else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
        if line is None: sh.line.fill.background()
        else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
        sh.shadow.inherit=False; return sh
    def label(x,y,w,h,lines,anchor=MSO_ANCHOR.MIDDLE,al=PP_ALIGN.CENTER):
        tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
        for m in ("left","right","top","bottom"): setattr(tf,f"margin_{m}",Pt(2))
        for i,(t,sz,c,b,it) in enumerate(lines):
            para(tf.paragraphs[0] if i==0 else tf.add_paragraph(),t,sz,c,b,it,al)
        return tb
    def text_in(shape,lines,anchor=MSO_ANCHOR.MIDDLE,al=PP_ALIGN.CENTER):
        tf=shape.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
        for m in ("left","right","top","bottom"): setattr(tf,f"margin_{m}",Pt(4))
        for i,(t,sz,c,b,it) in enumerate(lines):
            para(tf.paragraphs[0] if i==0 else tf.add_paragraph(),t,sz,c,b,it,al)
    def conn(x1,y1,x2,y2,color=CYAN,width=2.75,arrow=True):
        cn=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        cn.line.color.rgb=color; cn.line.width=Pt(width)
        ln=cn.line._get_or_add_ln()
        if arrow: ln.append(ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'lg'}))
        cn.shadow.inherit=False; return cn
    def monitor(cx,top,color=CYAN):
        w=0.95;h=0.62; box(cx-w/2,top,w,h,color,rad=0.12)
        box(cx-0.06,top+h,0.12,0.14,NAVY2,rad=0); box(cx-0.28,top+h+0.14,0.56,0.09,NAVY2,rad=0.3)
    def router_glyph(cx,top,color=CYAN):
        box(cx-0.5,top,1.0,0.5,color,rad=0.25); label(cx-0.5,top,1.0,0.5,[("⇄",14,WHITE,True,False)])

    cy=3.15; ch=4.35
    cx1=0.7; cw1=4.35
    box(cx1,cy,cw1,ch,WHITE,BORDER,1.5,0.05); monitor(cx1+cw1/2,cy+0.28)
    label(cx1,cy+1.15,cw1,0.55,[("MCP-Client (desktop)",13,NAVY,True,False),("172.18.0.2",10,MUTED,False,False)],MSO_ANCHOR.TOP)
    box(cx1+0.35,cy+2.05,cw1-0.7,1.95,NAVY,rad=0.06)
    label(cx1+0.5,cy+2.2,cw1-1.0,0.4,[("Host application",11.5,CYAN,True,False),("LibreChat",10,ICE,False,False)],MSO_ANCHOR.TOP)
    box(cx1+0.7,cy+3.1,cw1-1.4,0.75,CYAN,rad=0.18); label(cx1+0.7,cy+3.1,cw1-1.4,0.75,[("MCP Client",11,WHITE,True,False)])

    clx=5.55; clw=2.35; box(clx,cy+1.1,clw,1.9,TINT,CYAN_LT,1.5,shape=MSO_SHAPE.CLOUD)
    label(clx,cy+1.35,clw,1.4,[("Lab network",10.5,NAVY,True,False),("public-internet2",8.5,MUTED,False,True),("→ switch",8.5,MUTED,False,True)])

    sx=8.15; sw=5.15
    box(sx,cy,sw,ch,WHITE,BORDER,1.5,0.05); monitor(sx+sw/2,cy+0.28)
    label(sx,cy+1.15,sw,0.55,[("MCP-Server (desktop)",13,NAVY,True,False),("172.18.1.2",10,MUTED,False,False)],MSO_ANCHOR.TOP)
    box(sx+0.35,cy+2.05,sw-0.7,1.95,NAVY,rad=0.06)
    label(sx+0.5,cy+2.2,sw-1.0,0.5,[("MCP server script",11.5,CYAN,True,False),("FastMCP  ·  app/main.py",9.5,ICE,False,False)],MSO_ANCHOR.TOP)
    box(sx+0.7,cy+3.15,sw-1.4,0.72,GREEN,rad=0.18); label(sx+0.7,cy+3.15,sw-1.4,0.72,[("Tools:  get_interfaces · configure_interface",10,NAVY,True,False)])

    rx=13.9; rw=4.0
    box(rx,cy,rw,ch,WHITE,BORDER,1.5,0.05); router_glyph(rx+rw/2,cy+0.35)
    label(rx,cy+1.15,rw,0.55,[("Router1",13,NAVY,True,False),("172.18.2.6",10,MUTED,False,False)],MSO_ANCHOR.TOP)
    box(rx+0.4,cy+2.05,rw-0.8,1.95,TINT,BORDER,1.0,0.06)
    label(rx+0.4,cy+2.15,rw-0.8,1.8,[("Cisco IOS device",11,NAVY,True,False),("(the thing we actually",9.5,MUTED,False,True),("read & configure)",9.5,MUTED,False,True)],MSO_ANCHOR.TOP)

    midy=cy+ch/2
    conn(cx1+cw1,midy,clx,midy,CYAN,2.75); conn(clx+clw,midy,sx,midy,CYAN,2.75)
    label(cx1+cw1-0.2,midy-1.05,sx-cx1-cw1+0.4,0.9,[("MCP",11,CYAN,True,False),("Streamable HTTP",9,MUTED,False,False),(":8000/mcp · Bearer",9,MUTED,False,False)])
    conn(sx+sw,midy,rx,midy,GREEN,2.75)
    label(sx+sw-0.1,midy-0.95,rx-sx-sw+0.2,0.8,[("Netmiko SSH",10.5,GREEN_D,True,False),("show / config",9,MUTED,False,False)])

    foot=box(0.7,8.6,18.6,1.1,NAVY,rad=0.05)
    text_in(foot,[("One flow:  you ask LibreChat  →  its MCP client calls the server over MCP  →  the server's tools SSH into Router1  →  the answer comes back to you.",12.5,WHITE,True,False)])

    # move before 'Workshop lab topology'
    sldIdLst=prs.slides._sldIdLst
    new_el=list(sldIdLst)[-1]
    labidx=None
    for i,s in enumerate(prs.slides):
        try:
            if s.shapes.title and s.shapes.title.text.strip().startswith("Workshop lab topology"):
                labidx=i; break
        except Exception: pass
    sldIdLst.remove(new_el)
    if labidx is not None:
        sldIdLst.insert(list(sldIdLst).index(list(sldIdLst)[labidx]), new_el)
    else:
        sldIdLst.append(new_el)
    prs.save(path)
    print("added big-picture slide in", path)


if __name__=="__main__":
    main(sys.argv[1] if len(sys.argv)>1 else "Training Process.pptx")
